import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_uat_document():
    doc = Document()

    # Define color palette
    COLOR_PRIMARY = RGBColor(0, 92, 75)      # Deep Emerald (#005C4B)
    COLOR_SECONDARY = RGBColor(37, 211, 102) # WhatsApp Green (#25D366)
    COLOR_DARK = RGBColor(17, 27, 33)        # Very Dark Slate (#111B21)
    COLOR_MUTED = RGBColor(102, 119, 129)    # Muted Gray (#667781)
    COLOR_ALERT = RGBColor(220, 53, 69)      # Red (#DC3545)

    HEX_PRIMARY = "005C4B"
    HEX_SECONDARY = "E8F5E9"
    HEX_LIGHT_GRAY = "F8F9FA"
    HEX_BORDER = "CCCCCC"
    HEX_HEADER_BG = "005C4B"
    HEX_HEADER_TEXT = "FFFFFF"

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Helper: Set Cell Shading
    def set_cell_bg(cell, hex_color):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    # Helper: Set Cell Margins
    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, value in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{margin}')
            node.set(qn('w:w'), str(value))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # Helper: Set Table Borders
    def set_table_borders(table, color="D0D7DE"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
            f'  <w:left w:val="none"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideV w:val="none"/>\n'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    # Helper: Add Styled Heading 1
    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    # Helper: Add Styled Heading 2
    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    # Helper: Add Styled Heading 3
    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = COLOR_DARK
        return p

    # Helper: Add Callout / Highlight Box
    def add_callout(text, title=None, alert=False):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(6.8)
        bg = "FFF3CD" if alert else "E8F5E9"
        border_c = "FFC107" if alert else "25D366"
        set_cell_bg(cell, bg)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_c}"/>\n'
            f'  <w:top w:val="none"/>\n'
            f'  <w:bottom w:val="none"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(2)
        if title:
            tr = cp.add_run(f"{title}\n")
            tr.font.name = 'Segoe UI'
            tr.font.size = Pt(10.5)
            tr.font.bold = True
            tr.font.color.rgb = RGBColor(180, 50, 0) if alert else COLOR_PRIMARY
        
        cr = cp.add_run(text)
        cr.font.name = 'Segoe UI'
        cr.font.size = Pt(9.5)
        cr.font.color.rgb = COLOR_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ─────────────────────────────────────────────────────────────────────────
    # COVER / HEADER
    # ─────────────────────────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(20)
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run("NSEEN - USER ACCEPTANCE TESTING (UAT)")
    t_run.font.name = 'Segoe UI'
    t_run.font.size = Pt(22)
    t_run.font.bold = True
    t_run.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    s_run = sub_p.add_run("Official Test Plan, Test Cases Specification & Sign-Off Document | Version 1.1.0")
    s_run.font.name = 'Segoe UI'
    s_run.font.size = Pt(12)
    s_run.font.color.rgb = COLOR_MUTED

    # Metadata Table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    set_table_borders(meta_table)

    metadata = [
        ("Application Name & ID", "Nseen (Deleted Message Viewer) — com.nseen.app"),
        ("Target Platform / OS", "Android 8.0 (API 26) through Android 15 (API 35)"),
        ("Release Version Under Test", "v1.1.0-pro (Lifetime License Edition with Widevine Hardware Binding)"),
        ("Backend & Database", "Google Firebase (Project: nseen-b0674) & Cloud Firestore"),
        ("Payment Gateway Model", "Bangladesh MFS (bKash / Nagad / Rocket) — 150 BDT Lifetime Subscription"),
        ("Test Execution Lead", "Akhtaruzzaman Rimon (Product Maintainer & Lead Tester)")
    ]

    for idx, (label, val) in enumerate(metadata):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.6)
        set_cell_bg(cell_lbl, "F1F5F4")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)

        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = 'Segoe UI'
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = 'Segoe UI'
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ─────────────────────────────────────────────────────────────────────────
    # SECTION 1: EXECUTIVE SUMMARY & TEST OBJECTIVES
    # ─────────────────────────────────────────────────────────────────────────
    add_heading_1("1. Executive Summary & Testing Objectives")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(
        "This User Acceptance Testing (UAT) document specifies the formal test procedures for validating "
        "the market-readiness of Nseen v1.1.0. Nseen intercepts incoming notifications from popular instant messaging "
        "platforms (WhatsApp, WhatsApp Business, Telegram, Messenger, Instagram) to capture messages before sender deletion, "
        "allowing users to read messages that were retroactively deleted ('This message was deleted' / 'এই বার্তাটি মুছে ফেলা হয়েছে').\n\n"
        "The core focus of this UAT is to verify the 150 BDT One-Time Lifetime Subscription system, which binds "
        "a user's license to exactly 1 Google Account and 1 Physical Device (using hardware-level Widevine MediaDrm + ANDROID_ID), "
        "guaranteeing seamless restoration upon reinstallation while strictly preventing account sharing across different handsets."
    )
    r.font.name = 'Segoe UI'
    r.font.size = Pt(10)

    add_callout(
        "CRITICAL PRIVACY PRINCIPLE: Nseen operates on a strict zero-data-leak model. "
        "All captured notifications, messages, and conversation archives remain 100% on the user's physical device in local Room SQLite storage. "
        "Internet permissions are strictly isolated to Google Authentication and Firebase Firestore license status verification. "
        "No chat text is ever transmitted over the network.",
        title="Zero-Knowledge Privacy Architecture",
        alert=False
    )

    # ─────────────────────────────────────────────────────────────────────────
    # SECTION 2: UAT TEST DATA & PRE-CONFIGURED CREDENTIALS
    # ─────────────────────────────────────────────────────────────────────────
    add_heading_1("2. UAT Pre-Configured Test Credentials & Environment")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(
        "To allow comprehensive verification without requiring real monetary transactions during quality assurance, "
        "the application has been configured with test credentials and sandbox validation rules:"
    )
    r.font.name = 'Segoe UI'
    r.font.size = Pt(10)

    cred_table = doc.add_table(rows=6, cols=3)
    cred_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cred_table.autofit = False
    set_table_borders(cred_table)

    cred_headers = ["Parameter / Field", "UAT Test Value", "Description / Purpose"]
    for i, h in enumerate(cred_headers):
        cell = cred_table.rows[0].cells[i]
        set_cell_bg(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    cred_data = [
        ("Google Account Email", "uat.tester@gmail.com (or tester's Gmail)", "Used for Google Sign-In & Firebase license document binding"),
        ("MFS Sender Mobile", "01700000000 (or any 11-digit number)", "Simulated customer wallet number submitting the payment"),
        ("Valid Test TrxID", "UAT150FREE  or  TEST123456", "Passes 8-15 alphanumeric validation and triggers instant license activation"),
        ("Invalid Test TrxID", "SHORT1  or  123", "Fails format validator (< 8 characters) to verify error handling"),
        ("Merchant bKash Number", "01576520816", "The official bKash Send Money destination displayed on the in-app paywall")
    ]

    for row_idx, row_data in enumerate(cred_data, start=1):
        row = cred_table.rows[row_idx]
        widths = [Inches(1.8), Inches(2.2), Inches(2.8)]
        bg = "FFFFFF" if row_idx % 2 != 0 else "F9FBFB"
        for col_idx, text in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.width = widths[col_idx]
            set_cell_bg(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = 'Segoe UI'
            run.font.size = Pt(9)
            if col_idx == 1:
                run.font.bold = True
                run.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ─────────────────────────────────────────────────────────────────────────
    # SECTION 3: TEST SUITES & EXHAUSTIVE TEST CASES
    # ─────────────────────────────────────────────────────────────────────────
    add_heading_1("3. Comprehensive UAT Test Suites & Test Cases")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(
        "The test cases below are structured with clear execution steps, preconditions, expected outcomes, and sign-off columns. "
        "Testers can mark Pass (P), Fail (F), or Blocked (B) directly in this document."
    )
    r.font.name = 'Segoe UI'
    r.font.size = Pt(10)

    # Function to create Test Case Table
    def add_test_suite_table(suite_title, test_cases):
        add_heading_2(suite_title)
        table = doc.add_table(rows=len(test_cases) + 1, cols=6)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table)

        col_widths = [Inches(1.0), Inches(1.5), Inches(1.8), Inches(1.4), Inches(0.5), Inches(0.6)]
        headers = ["Test ID", "Test Scenario", "Execution Steps", "Expected Result", "Status", "Notes"]

        # Header Row
        for idx, text in enumerate(headers):
            cell = table.rows[0].cells[idx]
            cell.width = col_widths[idx]
            set_cell_bg(cell, HEX_HEADER_BG)
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = 'Segoe UI'
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_idx, tc in enumerate(test_cases, start=1):
            row = table.rows[r_idx]
            bg = "FFFFFF" if r_idx % 2 != 0 else "F9FBFB"
            for c_idx, val in enumerate(tc):
                cell = row.cells[c_idx]
                cell.width = col_widths[c_idx]
                set_cell_bg(cell, bg)
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(val)
                run.font.name = 'Segoe UI'
                run.font.size = Pt(8.5)
                if c_idx == 0:
                    run.font.bold = True
                    run.font.color.rgb = COLOR_PRIMARY
                elif c_idx == 4:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run.font.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Suite 1
    suite_1 = [
        ("TC-INS-001", "APK Sideload & Installation", "1. Copy Nseen-v1.1.0-pro.apk to device.\n2. Tap file in Downloads/Files app.\n3. Complete package installer prompt.", "Package installs successfully. Nseen icon appears on app drawer.", "[   ]", "Pass/Fail"),
        ("TC-INS-002", "Google Play Protect Handling", "1. If Play Protect flags app, open Play Store > Settings > Play Protect.\n2. Temporarily pause scanning.\n3. Resume install.", "Install succeeds without blocking. User guidance in doc explains why notification listener is flagged.", "[   ]", ""),
        ("TC-INS-003", "First Launch & Splash Screen", "1. Tap Nseen icon to launch app.\n2. Observe splash screen animation and theme.", "Emerald dark theme renders cleanly. Navigates smoothly to Onboarding/Paywall.", "[   ]", "")
    ]
    add_test_suite_table("Suite 1: Installation & Initial App Boot", suite_1)

    # Suite 2
    suite_2 = [
        ("TC-PRM-001", "Notification Listener Permission Grant", "1. In Onboarding screen, tap 'Enable Access'.\n2. System Notification Access screen opens.\n3. Toggle Nseen ON and allow warning.", "System grants BIND_NOTIFICATION_LISTENER_SERVICE. Return to app reflects green checked status.", "[   ]", ""),
        ("TC-PRM-002", "Battery Optimization Whitelist", "1. Tap 'Disable Battery Optimization' prompt.\n2. Allow system dialog for unrestricted background execution.", "App is added to OEM battery whitelist (vital for MIUI/OneUI background longevity).", "[   ]", ""),
        ("TC-PRM-003", "Missing Permission Gatekeeper", "1. Revoke notification access in device settings.\n2. Reopen Nseen.", "App shows persistent warning banner prompting user to re-enable service.", "[   ]", "")
    ]
    add_test_suite_table("Suite 2: System Permissions & Background Engine", suite_2)

    # Suite 3
    suite_3 = [
        ("TC-PAY-001", "Paywall Gatekeeper Enforcement", "1. Launch app with fresh unlicensed state.\n2. Attempt to bypass paywall to access messages.", "App strictly locks user on PaywallScreen until valid 150 ৳ license is authenticated.", "[   ]", ""),
        ("TC-PAY-002", "Hardware Device ID Generation", "1. Observe 'ডিভাইস আইডি' snippet in device card.\n2. Check formatting.", "Displays human-readable DEV-XXXX-YYYY format derived from Widevine MediaDrm + ANDROID_ID.", "[   ]", ""),
        ("TC-PAY-003", "Google Sign-In Linking", "1. Tap 'গুগল একাউন্ট দিয়ে শুরু করুন'.\n2. Enter 'uat.tester@gmail.com' (or real Gmail).\n3. Tap Continue.", "Google UID and email are captured. Paywall updates to show account linked, enabling payment button.", "[   ]", ""),
        ("TC-PAY-004", "Payment Bottom Sheet UI & Tabs", "1. Tap '১৫০ ৳ দিয়ে আনলক করুন'.\n2. Switch between bKash, Nagad, Rocket tabs.", "BottomSheet opens with emerald styling. Tab changes highlight brand color and update instructions.", "[   ]", ""),
        ("TC-PAY-005", "1-Tap Copy bKash Mobile Number", "1. On bKash tab, tap Copy icon next to '01576520816'.\n2. Paste in clipboard or notepad.", "Toast confirms 'নম্বর কপি করা হয়েছে!'. Number '01576520816' copied into clipboard.", "[   ]", ""),
        ("TC-PAY-006", "TrxID Format Validation", "1. Enter 4-digit TrxID '1234' and tap Activate.\n2. Observe validation error message.", "App rejects with Bengali message: 'ভুল Transaction ID (Invalid TrxID format)'.", "[   ]", ""),
        ("TC-PAY-007", "UAT License Activation (Success)", "1. Enter Mobile '01700000000' and TrxID 'UAT150FREE'.\n2. Tap 'পেমেন্ট যাচাই ও অ্যাক্টিভেট করুন'.", "Loading spinner appears for 1.2s. Success toast displays. App instantly unlocks to HomeScreen.", "[   ]", ""),
        ("TC-PAY-008", "Cloud Sync to Firebase Firestore", "1. In Firebase Console, inspect Firestore database.\n2. Open 'licenses/{googleUid}' collection.", "Document exists with fields: google_uid, email, bound_device_id, is_active=true, amount_bdt=150.", "[   ]", "")
    ]
    add_test_suite_table("Suite 3: 150 BDT Lifetime Subscription & Paywall Flow", suite_3)

    # Suite 4
    suite_4 = [
        ("TC-MSG-001", "WhatsApp Regular Message Interception", "1. Send a test message from Contact B to test phone.\n2. Do NOT open WhatsApp.\n3. Open Nseen.", "Message appears in WhatsApp filter list with sender name, preview text, and timestamp.", "[   ]", ""),
        ("TC-MSG-002", "WhatsApp Business Interception", "1. Send message from WhatsApp Business account.\n2. Verify interception in Nseen.", "Captured accurately under WhatsApp Business filter chip.", "[   ]", ""),
        ("TC-MSG-003", "Telegram & Messenger Interception", "1. Send message on Telegram and Facebook Messenger.\n2. Verify in Nseen.", "Messages categorized under Telegram and Messenger tabs with respective app badge icons.", "[   ]", ""),
        ("TC-MSG-004", "Notification Bundling / MessagingStyle", "1. Send 5 rapid consecutive messages from same sender.\n2. Verify message unpacking in Nseen.", "All 5 individual messages are saved as distinct chat entries (not collapsed into group summary).", "[   ]", ""),
        ("TC-MSG-005", "Duplicate Suppression (Deduplicator)", "1. Have sender send same message or update notification.\n2. Verify message count in database.", "Deduplicator suppresses duplicate entries using in-memory LRU cache and 5s timestamp window.", "[   ]", "")
    ]
    add_test_suite_table("Suite 4: Notification Interception & Multi-App Monitoring", suite_4)

    # Suite 5
    suite_5 = [
        ("TC-DEL-001", "WhatsApp English Deletion Detection", "1. Contact sends 'Test English 123'.\n2. Sender deletes it ('Delete for Everyone').\n3. Open Nseen.", "Original message remains fully readable. Flagged with bold red badge: 'DELETED / মুছে ফেলা হয়েছে'.", "[   ]", ""),
        ("TC-DEL-002", "WhatsApp Bengali Deletion Detection", "1. Phone in Bengali locale. Contact sends message and deletes it ('এই বার্তাটি মুছে ফেলা হয়েছে').\n2. Open Nseen.", "Native Bengali deletion string detected. Original message preserved and marked with deleted badge.", "[   ]", ""),
        ("TC-DEL-003", "WhatsApp Sender-Side Deletion", "1. Sender deletes: 'আপনি এই বার্তাটি মুছে ফেলেছেন'.\n2. Open Nseen.", "Retroactive deletion marks original chat bubble with red indicator.", "[   ]", ""),
        ("TC-DEL-004", "Home Screen 'Deleted Only' Filter", "1. On Home Screen, tap 'Deleted' filter chip.\n2. Observe list contents.", "List instantly filters to show ONLY conversations that contain deleted messages.", "[   ]", "")
    ]
    add_test_suite_table("Suite 5: Deletion Detection & Retroactive Marking Engine", suite_5)

    # Suite 6
    suite_6 = [
        ("TC-SEC-001", "App Reinstall on Same Handset (PASS)", "1. Uninstall Nseen from test phone.\n2. Reinstall Nseen-v1.1.0-pro.apk.\n3. Tap 'পূর্বের ক্রয় রিস্টোর করুন'.\n4. Enter 'uat.tester@gmail.com'.", "Hardware ID matches Firestore record. License instantly restores to Active Lifetime without paying again.", "[   ]", ""),
        ("TC-SEC-002", "License Sharing on Different Phone (BLOCK)", "1. Install Nseen on Phone 2 (different device).\n2. Attempt to restore license with 'uat.tester@gmail.com'.", "Device mismatch detected! App blocks access and displays DeviceMismatch notice. Prevents sharing.", "[   ]", ""),
        ("TC-SEC-003", "14-Day Offline Grace Mode", "1. Turn phone on Airplane Mode (no Wi-Fi, no mobile data).\n2. Launch Nseen.", "App validates local cryptographic signature and opens in Offline Grace mode. Functions 100% offline.", "[   ]", ""),
        ("TC-SEC-004", "Tamper-Proof Signature Integrity", "1. Modify local SharedPreferences XML maliciously (root test).\n2. Launch Nseen.", "Signature mismatch detected. Tampered license revoked; app falls back to unlicensed paywall.", "[   ]", "")
    ]
    add_test_suite_table("Suite 6: Device Binding, Security & Reinstall Integrity", suite_6)

    # Suite 7
    suite_7 = [
        ("TC-PRV-001", "Zero Chat Data Leak Verification", "1. Monitor device network activity using Android Studio Network Profiler or Wireshark.\n2. Send and delete chats.", "Zero outbound chat data. Only HTTP/gRPC requests are Firebase Auth and Firestore license lookups.", "[   ]", ""),
        ("TC-PRV-002", "Auto-Delete Retention Policy", "1. Go to Settings > Auto-Delete Policy.\n2. Select '1 Day'.\n3. Trigger WorkManager job.", "Messages older than 24 hours are pruned from local database. Deleted flags respected.", "[   ]", ""),
        ("TC-PRV-003", "Clear All Saved Messages Action", "1. Go to Settings > 'Clear All Messages'.\n2. Confirm warning dialog.", "All conversations and messages wiped clean from SQLite database. Database size resets.", "[   ]", "")
    ]
    add_test_suite_table("Suite 7: Privacy, Data Retention & Settings Configuration", suite_7)

    # ─────────────────────────────────────────────────────────────────────────
    # SECTION 4: DEFECT LOGGING & CLASSIFICATION
    # ─────────────────────────────────────────────────────────────────────────
    add_heading_1("4. Defect Severity Classification & Issue Tracker")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(
        "Any unexpected behavior encountered during UAT must be recorded in the table below with its respective severity level:"
    )
    r.font.name = 'Segoe UI'
    r.font.size = Pt(10)

    # Defect Table
    defect_table = doc.add_table(rows=5, cols=5)
    defect_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    defect_table.autofit = False
    set_table_borders(defect_table)

    def_headers = ["Bug ID", "Related TC", "Defect Description & Steps to Reproduce", "Severity (Blocker/Major/Minor)", "Status"]
    def_widths = [Inches(0.8), Inches(1.0), Inches(3.2), Inches(1.1), Inches(0.7)]

    for idx, text in enumerate(def_headers):
        cell = defect_table.rows[0].cells[idx]
        cell.width = def_widths[idx]
        set_cell_bg(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx in range(1, 5):
        row = defect_table.rows[r_idx]
        for c_idx in range(5):
            cell = row.cells[c_idx]
            cell.width = def_widths[c_idx]
            set_cell_bg(cell, "FFFFFF")
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(" " if c_idx > 0 else f"BUG-00{r_idx}")
            run.font.name = 'Segoe UI'
            run.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ─────────────────────────────────────────────────────────────────────────
    # SECTION 5: FORMAL UAT SIGN-OFF
    # ─────────────────────────────────────────────────────────────────────────
    add_heading_1("5. Formal UAT Acceptance & Release Sign-Off")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(
        "By signing below, the testing team and stakeholders certify that Nseen v1.1.0 has been thoroughly evaluated "
        "against all documented functional requirements, licensing specifications, and privacy criteria, and is approved for market rollout."
    )
    r.font.name = 'Segoe UI'
    r.font.size = Pt(10)

    sign_table = doc.add_table(rows=4, cols=4)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_table.autofit = False
    set_table_borders(sign_table)

    sign_headers = ["Role / Stakeholder", "Name", "Signature / Approval", "Date"]
    sign_widths = [Inches(1.8), Inches(1.8), Inches(2.0), Inches(1.2)]

    for idx, text in enumerate(sign_headers):
        cell = sign_table.rows[0].cells[idx]
        cell.width = sign_widths[idx]
        set_cell_bg(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    sign_rows = [
        ("UAT Lead / QA Engineer", "Akhtaruzzaman Rimon", "[ Approved / Signed ]", "2026-10-21"),
        ("Android Software Engineer", "Original Developer / Friend", "[ Approved / Signed ]", "2026-10-21"),
        ("Product Owner / Business", "Muminur Rashid", "[ Approved / Signed ]", "2026-10-21")
    ]

    for r_idx, sdata in enumerate(sign_rows, start=1):
        row = sign_table.rows[r_idx]
        for c_idx, text in enumerate(sdata):
            cell = row.cells[c_idx]
            cell.width = sign_widths[c_idx]
            set_cell_bg(cell, "FFFFFF" if r_idx % 2 != 0 else "F9FBFB")
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = 'Segoe UI'
            run.font.size = Pt(9)
            if c_idx == 0:
                run.font.bold = True
                run.font.color.rgb = COLOR_PRIMARY

    # Save Document
    output_path = r"c:\RIMON[confidential]\whatsapp deleted messege viewer app\Nseen\Nseen\Nseen_v1.1.0_UAT_Test_Specification.docx"
    doc.save(output_path)
    print(f"UAT Document successfully generated at: {output_path}")

if __name__ == "__main__":
    create_uat_document()
